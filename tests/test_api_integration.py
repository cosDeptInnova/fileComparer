from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.extractors import ExtractionResult
from app.main import app
from app.schemas import LLMComparisonResponse
from app.settings import settings


class StubLLMClient:
    model_name = "stub-llm"

    def compare(self, messages):
        user_payload = messages[-1]["content"]
        if '"rows"' in user_payload:
            return LLMComparisonResponse.model_validate({"changes": []})
        if "Texto nuevo añadido" in user_payload:
            return LLMComparisonResponse.model_validate(
                {
                    "changes": [
                        {
                            "change_type": "añadido",
                            "source_a": "",
                            "source_b": "Texto nuevo añadido",
                            "summary": "nuevo bloque",
                        }
                    ]
                }
            )
        return LLMComparisonResponse.model_validate(
            {
                "changes": [
                    {
                        "change_type": "modificado",
                        "source_a": "Texto base",
                        "source_b": "Texto base ajustado",
                        "summary": "cambio de redacción",
                    }
                ]
            }
        )


@pytest.fixture()
def client(monkeypatch, tmp_path: Path):
    from app.routes import comparar as comparar_route
    from app.services import jobs

    monkeypatch.setattr(settings, "inline_jobs", True)
    data_dir = tmp_path / "compare-jobs"
    data_dir.mkdir(parents=True, exist_ok=True)
    monkeypatch.setattr(settings, "data_dir", data_dir)

    state_store: dict[str, dict[str, object]] = {}

    def fake_update_job_state(job_id: str, **fields: object) -> None:
        state_store[job_id] = {**state_store.get(job_id, {}), "sid": job_id, **fields}

    def fake_read_job_state(sid: str) -> dict[str, object]:
        return dict(state_store.get(sid, {}))

    monkeypatch.setattr(comparar_route, "update_job_state", fake_update_job_state)
    monkeypatch.setattr(comparar_route, "read_job_state", fake_read_job_state)
    monkeypatch.setattr(jobs, "update_job_state", fake_update_job_state)
    return TestClient(app)


@pytest.fixture()
def csrf_token(client: TestClient) -> str:
    return client.get("/csrf-token").json()["csrf_token"]


def _post_compare(client: TestClient, csrf_token: str, *, files: dict[str, tuple[str, bytes, str]], data: dict[str, str] | None = None):
    return client.post(
        "/comparar",
        headers={"X-CSRFToken": csrf_token},
        cookies={settings.csrf_cookie_name: csrf_token},
        files=files,
        data=data or {},
    )


def test_compare_endpoint_contract(monkeypatch, client: TestClient, csrf_token: str, tmp_path: Path):
    from app.services import comparison_pipeline

    monkeypatch.setattr(comparison_pipeline, "LLMClient", lambda: StubLLMClient())

    file_a = tmp_path / "a.txt"
    file_b = tmp_path / "b.txt"
    file_a.write_text("Texto base", encoding="utf-8")
    file_b.write_text("Texto base ajustado. Texto nuevo añadido", encoding="utf-8")

    response = _post_compare(
        client,
        csrf_token,
        files={
            "file_a": ("a.txt", file_a.read_bytes(), "text/plain"),
            "file_b": ("b.txt", file_b.read_bytes(), "text/plain"),
        },
        data={"engine": "auto"},
    )
    assert response.status_code == 200
    sid = response.json()["sid"]

    progress = client.get(f"/progress/{sid}")
    assert progress.status_code == 200
    assert progress.json()["status"] == "done"

    result = client.get(f"/resultado/{sid}/json")
    assert result.status_code == 200
    payload = result.json()
    assert payload["sid"] == sid
    assert payload["rows"]
    assert payload["rows"][0]["change_type"] in {"añadido", "modificado", "eliminado"}
    assert payload["meta"]["documents"]["a"]["metadata"]["engine_used"] == "builtin"
    assert payload["meta"]["documents"]["a"]["metadata"]["conversion"]["applied"] is False


@pytest.mark.parametrize("extension", [".doc", ".xls", ".rtf"])
def test_legacy_formats_fail_without_soffice(monkeypatch, client: TestClient, csrf_token: str, extension: str):
    from app import extractors

    monkeypatch.setattr(extractors, "validate_soffice_option", lambda value: None)

    response = _post_compare(
        client,
        csrf_token,
        files={
            "file_a": (f"a{extension}", b"legacy-a", "application/octet-stream"),
            "file_b": (f"b{extension}", b"legacy-b", "application/octet-stream"),
        },
        data={"engine": "auto"},
    )

    assert response.status_code == 400
    assert "LibreOffice/soffice" in response.json()["detail"]


@pytest.mark.parametrize("extension", [".doc", ".xls", ".rtf"])
def test_legacy_formats_work_with_soffice(monkeypatch, client: TestClient, csrf_token: str, extension: str, tmp_path: Path):
    from app import extractors
    from app.services import comparison_pipeline

    monkeypatch.setattr(comparison_pipeline, "LLMClient", lambda: StubLLMClient())
    monkeypatch.setattr(extractors, "validate_soffice_option", lambda value: "/usr/bin/soffice" if value else None)

    def fake_convert_legacy(path: Path, soffice_path: str) -> Path:
        if path.suffix.lower() in {".doc", ".rtf"}:
            target = tmp_path / f"{path.stem}.docx"
            doc = Document()
            doc.add_paragraph(f"Contenido convertido {path.suffix.lower()}")
            doc.save(target)
            return target
        target = tmp_path / f"{path.stem}.xlsx"
        wb = Workbook()
        ws = wb.active
        ws["A1"] = f"Contenido convertido {path.suffix.lower()}"
        wb.save(target)
        return target

    monkeypatch.setattr(extractors, "_convert_legacy", fake_convert_legacy)

    response = _post_compare(
        client,
        csrf_token,
        files={
            "file_a": (f"a{extension}", b"legacy-a", "application/octet-stream"),
            "file_b": (f"b{extension}", b"legacy-b", "application/octet-stream"),
        },
        data={"engine": "auto", "soffice": "/usr/bin/soffice"},
    )

    assert response.status_code == 200
    sid = response.json()["sid"]
    payload = client.get(f"/resultado/{sid}/json").json()
    metadata_a = payload["meta"]["documents"]["a"]["metadata"]
    assert metadata_a["conversion"]["applied"] is True
    assert metadata_a["conversion"]["converter"] == "libreoffice"
    assert metadata_a["source_format_real"] == extension.lstrip(".")
    assert metadata_a["engine_used"] == "builtin"


def test_requested_engine_changes_extraction_result(monkeypatch, client: TestClient, csrf_token: str):
    from app.routes import comparar as comparar_route
    from app.services import comparison_pipeline

    monkeypatch.setattr(comparison_pipeline, "LLMClient", lambda: StubLLMClient())
    monkeypatch.setattr(comparar_route, "validate_extraction_request", lambda path, engine, soffice_path=None: (engine, None))

    def fake_extract_document_result(path: str, *, soffice_path=None, drop_headers=True, engine="auto"):
        resolved_engine = "builtin" if engine == "auto" else engine
        stem = Path(path).stem
        text = f"{resolved_engine}:{stem}"
        return ExtractionResult(
            text=text,
            engine=resolved_engine,
            quality_score=0.91,
            metadata={
                "source_format": "pdf",
                "source_format_real": "pdf",
                "conversion": {"applied": False},
                "engine_used": resolved_engine,
            },
            blocks=[],
            quality_signals={"block_count": 0},
        )

    monkeypatch.setattr(comparison_pipeline, "extract_document_result", fake_extract_document_result)

    common_files = {
        "file_a": ("a.pdf", b"%PDF-a", "application/pdf"),
        "file_b": ("b.pdf", b"%PDF-b", "application/pdf"),
    }

    builtin_response = _post_compare(
        client,
        csrf_token,
        files=common_files,
        data={"engine": "builtin"},
    )
    assert builtin_response.status_code == 200
    builtin_sid = builtin_response.json()["sid"]
    builtin_result = client.get(f"/resultado/{builtin_sid}/json").json()

    docling_response = _post_compare(
        client,
        csrf_token,
        files=common_files,
        data={"engine": "docling"},
    )
    assert docling_response.status_code == 200
    docling_sid = docling_response.json()["sid"]
    docling_result = client.get(f"/resultado/{docling_sid}/json").json()

    assert builtin_result["meta"]["documents"]["a"]["raw_text"] == "builtin:a"
    assert docling_result["meta"]["documents"]["a"]["raw_text"] == "docling:a"
    assert builtin_result["meta"]["documents"]["a"]["metadata"]["engine_used"] == "builtin"
    assert docling_result["meta"]["documents"]["a"]["metadata"]["engine_used"] == "docling"
