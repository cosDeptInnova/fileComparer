from pathlib import Path

from openpyxl import Workbook
from docx import Document

from app.services.extractors import extract_text_from_path


def test_extract_text_from_txt(tmp_path: Path):
    path = tmp_path / "sample.txt"
    path.write_text("hola mundo", encoding="utf-8")
    text, metadata = extract_text_from_path(path)
    assert text == "hola mundo"
    assert metadata["extension"] == ".txt"


def test_extract_text_from_docx(tmp_path: Path):
    path = tmp_path / "sample.docx"
    doc = Document()
    doc.add_paragraph("Linea uno")
    doc.save(path)
    text, metadata = extract_text_from_path(path)
    assert "Linea uno" in text
    assert metadata["extension"] == ".docx"


def test_extract_text_from_xlsx(tmp_path: Path):
    path = tmp_path / "sample.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "Hoja1"
    ws["A1"] = "Valor"
    wb.save(path)
    text, metadata = extract_text_from_path(path)
    assert "Valor" in text
    assert metadata["extension"] == ".xlsx"
